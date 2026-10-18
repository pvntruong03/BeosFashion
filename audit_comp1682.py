from docx import Document
from collections import Counter
from pathlib import Path
import re

path=r'C:\Users\pvntr\OneDrive\Desktop\Final 1\COMP-1682-001584160.docx'
d=Document(path)

def txt(p): return ' '.join(p.text.replace('\u202f',' ').replace('\xa0',' ').split())
def font_desc(r):
    return (r.font.name, r.font.size.pt if r.font.size else None, r.font.bold)

print('DOC',Path(path).name)
print('PARAGRAPHS',len(d.paragraphs),'TABLES',len(d.tables),'SECTIONS',len(d.sections))
for si,s in enumerate(d.sections,1):
 print('SECTION',si,'page_cm',round(s.page_width.cm,2),round(s.page_height.cm,2),'margins',round(s.top_margin.cm,2),round(s.bottom_margin.cm,2),round(s.left_margin.cm,2),round(s.right_margin.cm,2),'first_page_diff',s.different_first_page_header_footer,'footer_PAGE', ' PAGE ' in s.footer._element.xml)

# All non-empty paragraph direct properties, grouped.
items=[]
for i,p in enumerate(d.paragraphs):
 t=txt(p)
 if not t: continue
 runs=[font_desc(r) for r in p.runs if r.text.strip()]
 items.append((i,t,p.style.name,p.alignment,p.paragraph_format.line_spacing,runs))
print('STYLE_COUNT',Counter(x[2] for x in items))
print('ALIGN_COUNT',Counter(str(x[3]) for x in items))
print('LINE_SPACING',Counter(str(x[4]) for x in items))
print('RUN_FONTS',Counter(v for x in items for v in x[5]))

# Flag exceptions after Overview, outside intentional headings/title fields.
start=next(i for i,p in enumerate(d.paragraphs) if txt(p)=='Overview')
main={'Overview','Aim','Related Works / Literature Review','Objectives','Project Framework or Methodology Used','Legal, Social, Ethical and Professional','Planning (See Appendix A)','Declaration of AI Use'}
sub={'3.1 Related Works','3.2 Literature Review','System Requirements and Design','Customer Shopping System','Product and Store Management System','System Quality and Deployment','Legal','Social','Ethical','Professional'}
subsub=('3.1.1 ','3.1.2 ','3.1.3 ','3.1.4 ','3.1.5 ')
print('EXCEPTIONS')
for i,p in enumerate(d.paragraphs[start:],start):
 t=txt(p)
 if not t: continue
 runs=[r for r in p.runs if r.text.strip()]
 first=runs[0] if runs else None
 size=first.font.size.pt if first and first.font.size else None
 name=first.font.name if first else None
 bold=first.font.bold if first else None
 align=str(p.alignment)
 line=p.paragraph_format.line_spacing
 expected=None
 if t in main: expected=(16,True,'LEFT (0)')
 elif t in sub: expected=(14,True,'LEFT (0)')
 elif t.startswith(subsub): expected=(12,True,'LEFT (0)')
 elif t.startswith('Outcome describing this objective'): expected=(12,None,'LEFT (0)')
 elif t.startswith(('Analyse the functional','Design the system architecture','Develop customer registration','Implement size and colour selection','Develop product, category','Implement administrative','Test the completed','Deploy the completed')): expected=(12,True,'LEFT (0)')
 else: expected=(12,None,'JUSTIFY (3)')
 problems=[]
 if name not in ('Times New Roman',None): problems.append('font='+str(name))
 if size not in (expected[0],None): problems.append('size='+str(size))
 if expected[1] is True and bold is not True: problems.append('bold='+str(bold))
 if align != expected[2]: problems.append('align='+align)
 if line != 1.5: problems.append('line='+str(line))
 if problems: print(i,repr(t[:95]),'; '.join(problems))

print('TABLES')
for ti,table in enumerate(d.tables,1):
 allruns=[r for row in table.rows for cell in row.cells for p in cell.paragraphs for r in p.runs if r.text.strip()]
 fonts=Counter(font_desc(r) for r in allruns)
 lines=Counter(str(p.paragraph_format.line_spacing) for row in table.rows for cell in row.cells for p in cell.paragraphs if txt(p))
 print(ti,'rows',len(table.rows),'cols',len(table.columns),'fonts',fonts,'line',lines)
