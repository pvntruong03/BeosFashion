from docx import Document
from pathlib import Path
files=[r'C:\Users\pvntr\Downloads\3. Cover Page UoG (Sample).docx',r'C:\Users\pvntr\OneDrive\Desktop\Final 1\COMP-1682-001584160.docx',r'C:\Users\pvntr\OneDrive\Desktop\Final 1\COMP-1682-Final Year Project - citations standardised.docx',r'C:\Users\pvntr\OneDrive\Desktop\Final 1\COMP-1682-Final Year Project.docx',r'C:\Users\pvntr\OneDrive\Desktop\Final 1\COMP-1682-Final Year Project 3333.docx']
for f in files:
 try:
  d=Document(f); print('\n',Path(f).name, len(d.paragraphs),len(d.tables));
  for i,p in enumerate(d.paragraphs):
   if p.text.strip() and i<35:
    r=p.runs[0] if p.runs else None
    print(i, repr(p.text[:35]), p.style.name, r.font.name if r else None, r.font.size.pt if r and r.font.size else None, p.alignment)
 except Exception as e: print(f,e)
