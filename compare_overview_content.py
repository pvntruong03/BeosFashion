from docx import Document
from pathlib import Path
import re

A = r'C:\Users\pvntr\OneDrive\Desktop\Final 1\COMP-1682-001584160.docx'
B = r'C:\Users\pvntr\Downloads\3. Cover Page UoG (Sample) (1).docx'

def norm(s):
    return re.sub(r'\s+', ' ', s.replace('\u202f',' ').replace('\xa0',' ')).strip()

def content(path):
    d=Document(path)
    start=next(i for i,p in enumerate(d.paragraphs) if norm(p.text)=='Overview')
    paras=[norm(p.text) for p in d.paragraphs[start:] if norm(p.text)]
    tables=[]
    for t in d.tables:
        rows=[]
        for row in t.rows:
            rows.append(' | '.join(norm(c.text) for c in row.cells))
        tables.append(rows)
    return paras,tables
pa,ta=content(A); pb,tb=content(B)
print('PARAGRAPHS',len(pa),len(pb),'EQUAL',pa==pb)
maxlen=max(len(pa),len(pb))
for i in range(maxlen):
    a=pa[i] if i<len(pa) else '<MISSING>'
    b=pb[i] if i<len(pb) else '<MISSING>'
    if a!=b:
        print('DIFF',i+1)
        print(' A:',a)
        print(' B:',b)
print('TABLES',len(ta),len(tb),'EQUAL',ta==tb)
for i in range(max(len(ta),len(tb))):
    a=ta[i] if i<len(ta) else ['<MISSING>']
    b=tb[i] if i<len(tb) else ['<MISSING>']
    if a!=b:
        print('TABLE_DIFF',i+1)
        print(' A:',a)
        print(' B:',b)
