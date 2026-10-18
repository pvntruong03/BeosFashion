from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

path = r"C:\Users\pvntr\OneDrive\Desktop\Final 1\COMP-1682-001584160.docx"
doc = Document(path)

def clean(s):
    return ' '.join(s.replace('\u202f', ' ').replace('\xa0', ' ').split())

start = next(i for i, p in enumerate(doc.paragraphs) if clean(p.text) == 'Overview')
main_titles = {
    'Overview', 'Aim', 'Related Works / Literature Review', 'Objectives',
    'Project Framework or Methodology Used', 'Project Framework or Any Methodology used',
    'Legal, Social, Ethical and Professional', 'Planning (See Appendix A)',
    'Declaration of AI Use'
}
sub_titles = {
    '3.1 Related Works', '3.2 Literature Review',
    'System Requirements and Design', 'Customer Shopping System',
    'Product and Store Management System', 'System Quality and Deployment',
    'Introduction:', 'Legal', 'Social', 'Ethical', 'Professional'
}
objective_starts = (
    'Analyse the functional and non-functional requirements',
    'Design the system architecture, database structure',
    'Develop customer registration, login, and profile management',
    'Implement size and colour selection, shopping cart management',
    'Develop product, category, size, colour, product variant',
    'Implement administrative management, order management',
    'Test the completed system through unit testing',
    'Deploy the completed Fashion Store E-Commerce System'
)

def set_font(run, size=12, bold=None):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold

def format_paragraph(p, size=12, bold=None, alignment=None):
    p.paragraph_format.line_spacing = 1.5
    if alignment is not None:
        p.alignment = alignment
    for run in p.runs:
        set_font(run, size, bold)

# Apply format from Overview onwards; retain the title block before Overview.
for p in doc.paragraphs[start:]:
    text = clean(p.text)
    if text in main_titles:
        format_paragraph(p, 16, True, WD_ALIGN_PARAGRAPH.LEFT)
    elif text in sub_titles:
        format_paragraph(p, 14, True, WD_ALIGN_PARAGRAPH.LEFT)
    elif text.startswith(('3.1.1 ', '3.1.2 ', '3.1.3 ', '3.1.4 ', '3.1.5 ')) or text.startswith(objective_starts):
        format_paragraph(p, 12, True, WD_ALIGN_PARAGRAPH.LEFT)
    elif text.startswith('Outcome describing this objective'):
        format_paragraph(p, 12, True, WD_ALIGN_PARAGRAPH.LEFT)
    elif text == 'Table 1: Project Development Schedule':
        format_paragraph(p, 12, None, WD_ALIGN_PARAGRAPH.CENTER)
    else:
        format_paragraph(p, 12, None, WD_ALIGN_PARAGRAPH.JUSTIFY)

# Format the project development schedule only, leaving the cover-page table unchanged.
if len(doc.tables) >= 2:
    schedule = doc.tables[1]
    for row_index, row in enumerate(schedule.rows):
        for col_index, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index == 1 else WD_ALIGN_PARAGRAPH.CENTER
                format_paragraph(p, 12, row_index == 0, alignment)

# Page layout and centred page number. The first page stays unnumbered.
for section in doc.sections:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.5)
    section.right_margin = Cm(2)
    section.different_first_page_header_footer = True
    footer = section.footer
    for p in list(footer.paragraphs):
        p._element.getparent().remove(p._element)
    p = footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_font(run, 12)
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run._r.append(begin); run._r.append(instr); run._r.append(end)

doc.save(path)
print('SAVED', path)
