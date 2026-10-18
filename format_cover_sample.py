from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

src = r"C:\Users\pvntr\Downloads\3. Cover Page UoG (Sample).docx"
doc = Document(src)

# Page format and page numbering
for section in doc.sections:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.5)
    section.right_margin = Cm(2)
    section.different_first_page_header_footer = True
    footer = section.footer
    for p in footer.paragraphs:
        p._element.getparent().remove(p._element)
    p = footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.font.name = 'Times New Roman'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    r._r.append(begin); r._r.append(instr); r._r.append(end)

main_titles = {
    'Overview', 'Aim', 'Related Works / Literature Review', 'Objectives',
    'Project Framework or Methodology Used', 'Legal, Social, Ethical and Professional',
    'Planning (See Appendix A)', 'Declaration of AI Use\u202f'
}
sub_titles = {
    '3.1 Related Works', '3.2 Literature Review',
    'System Requirements and Design', 'Customer Shopping System',
    'Product and Store Management System', 'System Quality and Deployment',
    'Legal', 'Social', 'Ethical', 'Professional'
}
subsub_prefixes = ('3.1.1 ', '3.1.2 ', '3.1.3 ', '3.1.4 ', '3.1.5 ')

def set_font(run, size=12, bold=None):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold

def set_paragraph(p, size=12, bold=None, align=None, justify=False):
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align is not None:
        p.alignment = align
    for r in p.runs:
        set_font(r, size, bold)

# General formatting on all body/table text.
for p in doc.paragraphs:
    set_paragraph(p, 12, None, justify=(p.text.strip() != '' and p._p.getparent().tag.endswith('}body')))
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                set_paragraph(p, 12, None, justify=False)

# Preserve the centred title block on page 2.
for i in range(18, min(23, len(doc.paragraphs))):
    set_paragraph(doc.paragraphs[i], 12, None, WD_ALIGN_PARAGRAPH.CENTER)

# Heading hierarchy in the proposal text.
for p in doc.paragraphs:
    text = p.text.strip()
    if text in main_titles:
        set_paragraph(p, 16, True, WD_ALIGN_PARAGRAPH.LEFT)
    elif text in sub_titles:
        set_paragraph(p, 14, True, WD_ALIGN_PARAGRAPH.LEFT)
    elif text.startswith(subsub_prefixes):
        set_paragraph(p, 12, True, WD_ALIGN_PARAGRAPH.LEFT)

# Apply the declared style defaults.
for name in ('Normal', 'List Paragraph', 'Heading 1', 'Heading 2', 'Heading 3'):
    if name in doc.styles:
        st = doc.styles[name]
        st.font.name = 'Times New Roman'
        st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        st.font.size = Pt(12)

doc.save(src)
print('SAVED', src)
