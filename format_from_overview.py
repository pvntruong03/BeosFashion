from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

path = r"C:\Users\pvntr\Downloads\3. Cover Page UoG (Sample) (1).docx"
doc = Document(path)

start = next(i for i,p in enumerate(doc.paragraphs) if p.text.strip() == 'Overview')
main_titles = {
    'Overview', 'Aim', 'Related Works / Literature Review', 'Objectives',
    'Project Framework or Methodology Used', 'Legal, Social, Ethical and Professional',
    'Planning (See Appendix A)', 'Declaration of AI Use\u202f'
}
sub_titles = {
    '3.1 Related Works', '3.2 Literature Review',
    'System Requirements and Design', 'Customer Shopping System',
    'Product and Store Management System', 'System Quality and Deployment',
    'Legal', 'Social', 'Ethical', 'Professional'
}
objective_actions = {
    'Analyse the functional and non-functional requirements of the Fashion Store E-Commerce System, including customer accounts, product browsing, product filtering, shopping carts, checkout, orders, payments, inventory, and administration. [0.5 month]',
    'Design the system architecture, database structure, user interfaces, and system diagrams based on the approved requirements. The design will cover all functions included in the project scope. [0.5 month]',
    'Develop customer registration, login, and profile management functions. Enable customers to browse products, view product details, search by keyword, and filter products by gender, size, colour, price range, and collection. [1.0 month]',
    'Implement size and colour selection, shopping cart management, checkout, online payment options, order placement, and order history functions. [1.0 month]',
    'Develop product, category, size, colour, product variant, image, and inventory management functions. Stock quantities will be updated when administrators make inventory changes or when orders are processed. [1.0 month]',
    'Implement administrative management, order management, and role-based access control using Spring Security so that administrative functions are available only to authorised users. [0.5 month]',
    'Test the completed system through unit testing, integration testing, functional testing, and system testing. Record and correct identified defects before deployment. [1.0 month]',
    'Deploy the completed Fashion Store E-Commerce System to the target environment and verify customer accounts, product browsing and filtering, shopping carts, checkout, orders, payments, inventory, and administrative functions. [0.5 month]'
}

def font_run(run, size=12, bold=None):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold

def apply(p, size=12, bold=None, align=None):
    p.paragraph_format.line_spacing = 1.5
    if align is not None:
        p.alignment = align
    for run in p.runs:
        font_run(run, size, bold)

# Only format the proposal content beginning with Overview.
for p in doc.paragraphs[start:]:
    text = p.text.strip()
    if text in main_titles:
        apply(p, 16, True, WD_ALIGN_PARAGRAPH.LEFT)
    elif text in sub_titles:
        apply(p, 14, True, WD_ALIGN_PARAGRAPH.LEFT)
    elif text.startswith(('3.1.1 ', '3.1.2 ', '3.1.3 ', '3.1.4 ', '3.1.5 ')) or text in objective_actions:
        apply(p, 12, True, WD_ALIGN_PARAGRAPH.LEFT)
    else:
        apply(p, 12, None, WD_ALIGN_PARAGRAPH.JUSTIFY)

# The project schedule table is part of the proposal content.
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                apply(p, 12, None, WD_ALIGN_PARAGRAPH.LEFT)

# Apply page setup and a centred page-number field; first page remains unnumbered.
for section in doc.sections:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.5)
    section.right_margin = Cm(2)
    section.different_first_page_header_footer = True
    footer = section.footer
    for p in list(footer.paragraphs):
        p._element.getparent().remove(p._element)
    p = footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    font_run(run, 12)
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run._r.append(begin); run._r.append(instr); run._r.append(end)

doc.save(path)
print('SAVED', path)
