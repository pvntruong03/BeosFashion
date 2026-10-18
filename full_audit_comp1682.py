from docx import Document
from collections import Counter
from pathlib import Path
import re, zipfile

path=r'C:\Users\pvntr\OneDrive\Desktop\Final 1\COMP-1682-001584160.docx'
d=Document(path)

def clean(s): return ' '.join(s.replace('\u202f',' ').replace('\xa0',' ').split())
print('=== STRUCTURE ===')
print('paragraphs',len(d.paragraphs),'tables',len(d.tables),'sections',len(d.sections),'inline_shapes',len(d.inline_shapes))
for i,p in enumerate(d.paragraphs):
 t=clean(p.text)
 if t:
  print(f'P{i}: {t[:170]}')
print('=== PAGE_SETUP ===')
for i,s in enumerate(d.sections,1):
 print(i, 'A4',round(s.page_width.cm,2),round(s.page_height.cm,2),'margins',round(s.top_margin.cm,2),round(s.bottom_margin.cm,2),round(s.left_margin.cm,2),round(s.right_margin.cm,2),'first-page-diff',s.different_first_page_header_footer,'pagefield',(' PAGE ' in s.footer._element.xml))
print('=== FORMATTING ===')
all_runs=[r for p in d.paragraphs for r in p.runs if r.text.strip()]
all_runs += [r for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs if r.text.strip()]
print('fonts',Counter(r.font.name or '<inherit>' for r in all_runs))
print('sizes',Counter(r.font.size.pt if r.font.size else '<inherit>' for r in all_runs))
print('body_alignment',Counter(str(p.alignment) for p in d.paragraphs if clean(p.text)))
print('line_spacing',Counter(str(p.paragraph_format.line_spacing) for p in d.paragraphs if clean(p.text)))
print('=== CITATIONS ===')
text='\n'.join(clean(p.text) for p in d.paragraphs)
cites=re.findall(r'\(([A-Z][A-Za-z& ]+(?:et al\.)?,\s*(?:19|20)\d{2}[a-z]?)\)',text)
print('citations',cites)
for key in ['References','Reference','Bibliography','Appendix','Appendices','Use Case','Figure','Table 1']:
 print(key, bool(re.search(r'\b'+re.escape(key)+r'\b',text,re.I)))
print('=== TABLES ===')
for ti,t in enumerate(d.tables,1):
 print('TABLE',ti,'rows',len(t.rows),'cols',len(t.columns))
 for row in t.rows:
  print(' | '.join(clean(c.text) for c in row.cells))
print('=== PACKAGE ===')
with zipfile.ZipFile(path) as z:
 print('zip_ok',z.testzip() is None,'has_comments',any(n.startswith('word/comments') for n in z.namelist()),'has_footnotes',any(n=='word/footnotes.xml' for n in z.namelist()))
