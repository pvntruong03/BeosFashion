from docx import Document
from copy import deepcopy

out_path = r"C:\Users\pvntr\Downloads\3. Cover Page UoG (Sample).docx"
reference_path = r"C:\Users\pvntr\OneDrive\Desktop\Final 1\COMP-1682-001584160.docx"
out = Document(out_path)
ref = Document(reference_path)

def norm(s):
    return ' '.join(s.replace('\u202f',' ').replace('\xa0',' ').split())

ref_by_text = {}
for p in ref.paragraphs:
    key = norm(p.text)
    if key and key not in ref_by_text:
        ref_by_text[key] = p

matched = 0
for p in out.paragraphs:
    source = ref_by_text.get(norm(p.text))
    if source is None:
        continue
    old_ppr = p._p.pPr
    if old_ppr is not None:
        p._p.remove(old_ppr)
    if source._p.pPr is not None:
        p._p.insert(0, deepcopy(source._p.pPr))
    src_runs = {norm(r.text): r for r in source.runs if norm(r.text)}
    for r in p.runs:
        sr = src_runs.get(norm(r.text))
        if sr is None:
            continue
        old_rpr = r._r.rPr
        if old_rpr is not None:
            r._r.remove(old_rpr)
        if sr._r.rPr is not None:
            r._r.insert(0, deepcopy(sr._r.rPr))
    matched += 1

# Restore page geometry from the corresponding saved proposal without copying
# header/footer relationship identifiers between documents.
os, rs = out.sections[0], ref.sections[0]
os.page_width, os.page_height = rs.page_width, rs.page_height
os.top_margin, os.bottom_margin = rs.top_margin, rs.bottom_margin
os.left_margin, os.right_margin = rs.left_margin, rs.right_margin
os.header_distance, os.footer_distance = rs.header_distance, rs.footer_distance
os.gutter = rs.gutter
os.different_first_page_header_footer = rs.different_first_page_header_footer

out.save(out_path)
print('RESTORED_MATCHED_PARAGRAPHS', matched)
