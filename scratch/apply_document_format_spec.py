from copy import deepcopy
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

FILE = Path(r"C:\Users\pvntr\OneDrive\Desktop\Final 1\COMP-1682-Final Year Project - citations standardised.docx")
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
FONT_NAME = "Times New Roman"

SECTION_HEADINGS = {
    "Overview", "Aim", "Related Works / Literature Review", "Objectives",
    "Project Framework or Methodology Used", "Legal, Social, Ethical and Professional",
    "Planning (See Appendix A)", "Initial References", "Declaration of AI Use",
}
SUBSECTION_HEADINGS = {
    "3.1 Related Works", "3.2 Literature Review", "System Requirements and Design",
    "Customer Shopping System", "Product and Store Management System",
    "System Quality and Deployment", "Introduction:", "Legal", "Social", "Ethical",
    "Professional",
}


def set_run_font(run, size, bold=None):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = properties._add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), FONT_NAME)


def set_paragraph_format(paragraph, alignment=None):
    if alignment is not None:
        paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing = 1.5


def clear_paragraph(paragraph):
    properties = paragraph._element.pPr
    for child in list(paragraph._element):
        if child is not properties:
            paragraph._element.remove(child)


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, 12)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


document = Document(FILE)

# Exact A4 page size and margins.
for section in document.sections:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.5)
    section.right_margin = Cm(2)
    section.different_first_page_header_footer = True

    # First-page footer remains blank; remaining pages show a centred number.
    for paragraph in section.first_page_footer.paragraphs:
        clear_paragraph(paragraph)
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1.5
    add_page_field(paragraph)

# Page two begins with Overview, leaving the first page as the cover page.
overview = next((p for p in document.paragraphs if p.text.strip() == "Overview"), None)
if overview is not None:
    overview.paragraph_format.page_break_before = True

# Set normal and heading style defaults.
normal = document.styles["Normal"]
normal.font.name = FONT_NAME
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.5
for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
    style = document.styles[name]
    style.font.name = FONT_NAME
    style.font.size = Pt(size)
    style.font.bold = True
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

for paragraph in document.paragraphs:
    text = paragraph.text.strip()
    if text in SECTION_HEADINGS:
        set_paragraph_format(paragraph, WD_ALIGN_PARAGRAPH.LEFT)
        paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            set_run_font(run, 16, True)
    elif text in SUBSECTION_HEADINGS:
        set_paragraph_format(paragraph, WD_ALIGN_PARAGRAPH.LEFT)
        paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            set_run_font(run, 14, True)
    elif text.startswith("3.1."):
        set_paragraph_format(paragraph, WD_ALIGN_PARAGRAPH.LEFT)
        paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            set_run_font(run, 12, True)
    else:
        set_paragraph_format(paragraph)
        # Cover-page content, references, and captions remain centred or left;
        # the remaining narrative is justified.
        if text and text not in {"COMP1682.1 Project Proposal", "Development of a Fashion Store E-Commerce System Using Java and Spring Boot", "[Pham Vo Nhut Truong]", "BSc (Hons) Computing", "[001584160]", "Table 1: Project Development Schedule", "Use Case Diagram", "References"} and not text.startswith(("Keywords:", "Outcome describing", "To develop research", "To create an outline", "To explain concepts", "To support my use", "To summarise")):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in paragraph.runs:
            set_run_font(run, 12)

for table in document.tables:
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                set_paragraph_format(paragraph)
                if row_index > 0 and column_index == 1:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                for run in paragraph.runs:
                    set_run_font(run, 12, row_index == 0)

# Ensure Word updates the PAGE field when the document is opened.
settings = document.settings.element
update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set(qn("w:val"), "true")

with NamedTemporaryFile(delete=False, suffix=".docx", dir=FILE.parent) as temporary:
    document_temp = Path(temporary.name)
document.save(document_temp)

# Apply 1.5-line spacing to all document paragraphs, including References
# stored inside Word content controls that python-docx does not expose.
with ZipFile(document_temp, "r") as source:
    root = etree.fromstring(source.read("word/document.xml"))
    for paragraph in root.xpath("//w:p", namespaces={"w": W}):
        paragraph_properties = paragraph.find(f"{{{W}}}pPr")
        if paragraph_properties is None:
            paragraph_properties = etree.Element(f"{{{W}}}pPr")
            paragraph.insert(0, paragraph_properties)
        spacing = paragraph_properties.find(f"{{{W}}}spacing")
        if spacing is None:
            spacing = etree.SubElement(paragraph_properties, f"{{{W}}}spacing")
        spacing.set(f"{{{W}}}line", "360")
        spacing.set(f"{{{W}}}lineRule", "auto")
    updated_xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with NamedTemporaryFile(delete=False, suffix=".docx", dir=FILE.parent) as final_temp_file:
        final_temp = Path(final_temp_file.name)
    try:
        with ZipFile(final_temp, "w", ZIP_DEFLATED) as output:
            for item in source.infolist():
                output.writestr(item, updated_xml if item.filename == "word/document.xml" else source.read(item.filename))
        backup = Path(r"D:\web code\scratch\COMP-1682-Final Year Project - citations standardised.before-format-spec.docx")
        backup.write_bytes(FILE.read_bytes())
        with final_temp.open("rb") as incoming, FILE.open("r+b") as destination:
            destination.seek(0)
            destination.truncate(0)
            destination.write(incoming.read())
            destination.flush()
    finally:
        if final_temp.exists():
            final_temp.unlink()
        if document_temp.exists():
            document_temp.unlink()

print("Applied page, typography, heading, spacing, alignment, and page-number format.")
