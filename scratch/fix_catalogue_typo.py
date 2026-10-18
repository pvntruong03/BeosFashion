from pathlib import Path

from docx import Document

FILE = Path(r"C:\Users\pvntr\OneDrive\Desktop\Final 1\COMP-1682-Final Year Project - citations standardised.docx")


def fix_paragraph(paragraph):
    changed = 0
    for run in paragraph.runs:
        corrected = run.text.replace("catalogueue", "catalogue")
        if corrected != run.text:
            run.text = corrected
            changed += 1
    return changed


document = Document(FILE)
changed = sum(fix_paragraph(paragraph) for paragraph in document.paragraphs)
for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            changed += sum(fix_paragraph(paragraph) for paragraph in cell.paragraphs)

backup = Path(r"D:\web code\scratch\COMP-1682-Final Year Project - citations standardised.before-catalogue-fix.docx")
backup.write_bytes(FILE.read_bytes())
document.save(FILE)
print(f"Corrected {changed} runs.")
