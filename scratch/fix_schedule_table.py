from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


FILE = Path(r"C:\Users\pvntr\OneDrive\Desktop\Final 1\COMP-1682-Final Year Project - citations standardised.docx")
CAPTION = "Table 1: Project Development Schedule"
DURATIONS = ["2 weeks", "2 weeks", "14 weeks", "4 weeks", "1 week", "1 week"]


def set_repeat_table_header(row):
    row_properties = row._tr.get_or_add_trPr()
    existing = row_properties.find(qn("w:tblHeader"))
    if existing is None:
        existing = OxmlElement("w:tblHeader")
        row_properties.append(existing)
    existing.set(qn("w:val"), "true")


document = Document(FILE)
if len(document.tables) != 1:
    raise ValueError("Expected exactly one project schedule table.")

table = document.tables[0]
if len(table.rows) != 7 or len(table.columns) != 5:
    raise ValueError("The project schedule table structure is not as expected.")

for row, duration in zip(table.rows[1:], DURATIONS):
    cell = row.cells[4]
    cell.paragraphs[0].text = duration

set_repeat_table_header(table.rows[0])

has_caption = any(CAPTION in paragraph.text for paragraph in document.paragraphs)
if not has_caption:
    caption_xml = OxmlElement("w:p")
    paragraph_properties = OxmlElement("w:pPr")
    alignment = OxmlElement("w:jc")
    alignment.set(qn("w:val"), "center")
    paragraph_properties.append(alignment)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "120")
    paragraph_properties.append(spacing)
    caption_xml.append(paragraph_properties)
    table._tbl.addprevious(caption_xml)
    caption_paragraph = Paragraph(caption_xml, table._parent)
    caption_run = caption_paragraph.add_run(CAPTION)
    caption_run.font.bold = False

document.save(FILE)
