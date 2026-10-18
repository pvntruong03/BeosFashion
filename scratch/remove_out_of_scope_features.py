from pathlib import Path

from docx import Document

FILE = Path(r"C:\Users\pvntr\OneDrive\Desktop\Final 1\COMP-1682-Final Year Project - citations standardised.docx")

REPLACEMENTS = {
    "including diverse sizes, colourways, and material options": "including diverse sizes and colourways",
    "such as sizes, colours, and materials": "such as sizes and colours",
    "In addition, the system will support warehouse stock and vendor management, enabling administrators to manage items stored in the warehouse and coordinate supplier details. ": "In addition, the system will support warehouse stock management, enabling administrators to maintain accurate stock information. ",
    "To support business operations, the system will provide order management and sales reporting features, helping administrators track revenue, order fulfilment status, and overall performance.": "To support business operations, the system will provide order management features, helping administrators track order fulfilment status.",
    "Develop product, category, brand, size, colour, material, product variant, image, and inventory management functions.": "Develop product, category, size, colour, product variant, image, and inventory management functions.",
    "The design also defines product variants such as size, colour, material, price, and stock quantity.": "The design also defines product variants such as size, colour, price, and stock quantity.",
}


def update_paragraph(paragraph):
    changed = 0
    for run in paragraph.runs:
        value = run.text
        for old, new in REPLACEMENTS.items():
            value = value.replace(old, new)
        if value != run.text:
            run.text = value
            changed += 1
    return changed


document = Document(FILE)
changed = sum(update_paragraph(paragraph) for paragraph in document.paragraphs)
for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            changed += sum(update_paragraph(paragraph) for paragraph in cell.paragraphs)

# This paragraph contains several Word runs, so update its complete text to
# remove the final out-of-scope vendor sentence without leaving fragments.
aim = next((p for p in document.paragraphs if p.text.startswith("The aim of this project is to develop a fashion store e-commerce system")), None)
if aim is None:
    raise ValueError("Aim paragraph not found.")
updated_aim = aim.text.replace(
    "In addition, the system will support warehouse stock and vendor management, enabling administrators to manage items stored in the warehouse and coordinate supplier details. ",
    "In addition, the system will support warehouse stock management, enabling administrators to maintain accurate stock information. "
).replace("finalize orders", "finalise orders").replace("minimize warehouse", "minimise warehouse")
if updated_aim != aim.text:
    first_run = aim.runs[0]
    first_run.text = updated_aim
    for run in aim.runs[1:]:
        run.text = ""
    changed += 1

backup = Path(r"D:\web code\scratch\COMP-1682-Final Year Project - citations standardised.before-scope-alignment.docx")
backup.write_bytes(FILE.read_bytes())
document.save(FILE)
print(f"Removed out-of-scope features from {changed} runs.")
