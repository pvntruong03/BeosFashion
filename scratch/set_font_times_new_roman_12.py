from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

FILE = Path(r"C:\Users\pvntr\OneDrive\Desktop\Final 1\COMP-1682-Final Year Project - citations standardised.docx")
FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(12)


def set_run_font(run):
    run.font.name = FONT_NAME
    run.font.size = FONT_SIZE
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = properties._add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), FONT_NAME)


def format_paragraph(paragraph):
    for run in paragraph.runs:
        set_run_font(run)


document = Document(FILE)
changed = 0

for style in document.styles:
    if style.type == 1:  # paragraph style
        style.font.name = FONT_NAME
        style.font.size = FONT_SIZE

for paragraph in document.paragraphs:
    format_paragraph(paragraph)
    changed += len(paragraph.runs)

for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                format_paragraph(paragraph)
                changed += len(paragraph.runs)

for section in document.sections:
    for paragraph in list(section.header.paragraphs) + list(section.footer.paragraphs):
        format_paragraph(paragraph)
        changed += len(paragraph.runs)

backup = Path(r"D:\web code\scratch\COMP-1682-Final Year Project - citations standardised.before-font.docx")
backup.write_bytes(FILE.read_bytes())
document.save(FILE)
print(f"Applied {FONT_NAME} {FONT_SIZE.pt:g} pt to {changed} runs.")
