from pathlib import Path

from docx import Document

FILE = Path(r"C:\Users\pvntr\OneDrive\Desktop\Final 1\COMP-1682-Final Year Project - citations standardised.docx")

# Apply the project name first, then UK spelling across the document.
REPLACEMENTS = [
    ("BEO Fashion E-Commerce Website", "Fashion Store E-Commerce System"),
    ("BEO Fashion E-Commerce System", "Fashion Store E-Commerce System"),
    ("BEO Fashion system", "Fashion Store E-Commerce System"),
    ("Fashion Store E-Commerce Website", "Fashion Store E-Commerce System"),
    ("Fashion Store E-Commerce website", "Fashion Store E-Commerce System"),
    ("fashion store e-commerce system", "fashion store e-commerce system"),
    ("colors", "colours"),
    ("Colors", "Colours"),
    ("color", "colour"),
    ("Color", "Colour"),
    ("catalogs", "catalogues"),
    ("catalog", "catalogue"),
    ("Catalog", "Catalogue"),
    ("analyze", "analyse"),
    ("Analyze", "Analyse"),
    ("analyzed", "analysed"),
    ("Analyzed", "Analysed"),
    ("personalized", "personalised"),
    ("Personalized", "Personalised"),
    ("centralized", "centralised"),
    ("Centralized", "Centralised"),
    ("finalized", "finalised"),
    ("Finalized", "Finalised"),
    ("fulfillment", "fulfilment"),
    ("Fulfillment", "Fulfilment"),
]

HEADINGS = {
    "Develop a Fashion Store E-Commerce System using Java(....)": "Development of a Fashion Store E-Commerce System Using Java and Spring Boot",
    "Related works / Literature review": "Related Works / Literature Review",
    "3.1 Related works": "3.1 Related Works",
    "3.2 Literature review": "3.2 Literature Review",
    "Project Framework or Any Methodology used": "Project Framework or Methodology Used",
    "Planning (see appendix A)": "Planning (See Appendix A)",
    "Use-case Diagram": "Use Case Diagram",
}


def update_paragraph(paragraph):
    text = paragraph.text
    if text in HEADINGS:
        for run in paragraph.runs:
            run.text = ""
        if paragraph.runs:
            paragraph.runs[0].text = HEADINGS[text]
        return True
    changed = False
    for run in paragraph.runs:
        value = run.text
        for old, new in REPLACEMENTS:
            value = value.replace(old, new)
        if value != run.text:
            run.text = value
            changed = True
    return changed


document = Document(FILE)
changed = 0
for paragraph in document.paragraphs:
    changed += int(update_paragraph(paragraph))
for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                changed += int(update_paragraph(paragraph))
for section in document.sections:
    for paragraph in list(section.header.paragraphs) + list(section.footer.paragraphs):
        changed += int(update_paragraph(paragraph))

backup = Path(r"D:\web code\scratch\COMP-1682-Final Year Project - citations standardised.before-consistency.docx")
backup.write_bytes(FILE.read_bytes())
document.save(FILE)
print(f"Standardised {changed} paragraphs.")
