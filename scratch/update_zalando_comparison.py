from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


SOURCE = r"C:\Users\pvntr\OneDrive\Desktop\Final 1\COMP-1682-Final Year Project.docx"
TARGET_START = "A similar system that can be used as a reference is Zalando"

FIRST_PARAGRAPH = (
    "Zalando is used as a reference because it is a large multi-brand fashion "
    "e-commerce platform operating across different European markets. It provides "
    "customers with a wide range of fashion products, personalised recommendations, "
    "size guidance, and services for retail partners (Zalando, 2026). In comparison, "
    "the BEO Fashion E-Commerce Website will have a smaller and more focused scope. "
    "Rather than supporting many independent brands and sellers, it will manage the "
    "BEO Fashion product catalogue and its store operations. The system will include "
    "product search and filtering, product variants, shopping carts, online payment "
    "options, order history, and inventory management."
)

SECOND_PARAGRAPH = (
    "Unlike Zalando, the system will not operate as a marketplace for multiple "
    "independent brands or sellers. Instead, it will support BEO Fashion staff in "
    "managing products, stock, customers, orders, and store operations. The system "
    "will also include features such as store location, staff management, stock "
    "adjustment, and stock transfer between stores. Although its scope will be smaller "
    "than Zalando, it will focus on the functions needed for BEO Fashion to operate an "
    "online fashion store and manage its internal retail activities."
)


def copy_run_format(source_run, target_run):
    if source_run._r.rPr is not None:
        target_run._r.insert(0, deepcopy(source_run._r.rPr))


document = Document(SOURCE)
target = next((p for p in document.paragraphs if p.text.startswith(TARGET_START)), None)
if target is None:
    raise ValueError("The Zalando comparison paragraph was not found.")

source_run = target.runs[0] if target.runs else None
target.clear()
first_run = target.add_run(FIRST_PARAGRAPH)
if source_run is not None:
    copy_run_format(source_run, first_run)

new_element = OxmlElement("w:p")
if target._p.pPr is not None:
    new_element.append(deepcopy(target._p.pPr))
target._p.addnext(new_element)
second = Paragraph(new_element, target._parent)
second_run = second.add_run(SECOND_PARAGRAPH)
if source_run is not None:
    copy_run_format(source_run, second_run)

document.save(SOURCE)
