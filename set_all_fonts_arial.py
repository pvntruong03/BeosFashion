from docx import Document
from docx.oxml.ns import qn

path = r"C:\Users\pvntr\OneDrive\Desktop\Final 1\COMP-1682-001584160.docx"
doc = Document(path)

def set_arial(run):
    run.font.name = 'Arial'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    run._element.rPr.rFonts.set(qn('w:cs'), 'Arial')

def walk_paragraphs(paragraphs):
    for p in paragraphs:
        for r in p.runs:
            set_arial(r)

def walk_table(table):
    for row in table.rows:
        for cell in row.cells:
            walk_paragraphs(cell.paragraphs)
            for nested in cell.tables:
                walk_table(nested)

walk_paragraphs(doc.paragraphs)
for table in doc.tables:
    walk_table(table)
for section in doc.sections:
    walk_paragraphs(section.header.paragraphs)
    walk_paragraphs(section.first_page_header.paragraphs)
    walk_paragraphs(section.footer.paragraphs)
    walk_paragraphs(section.first_page_footer.paragraphs)
    for table in section.header.tables + section.first_page_header.tables + section.footer.tables + section.first_page_footer.tables:
        walk_table(table)
for style in doc.styles:
    if not hasattr(style, 'font'):
        continue
    style.font.name = 'Arial'
    style._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    style._element.rPr.rFonts.set(qn('w:cs'), 'Arial')

doc.save(path)
print('SAVED', path)

