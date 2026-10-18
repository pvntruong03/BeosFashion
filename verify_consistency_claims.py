from docx import Document
import re
p=r'C:\Users\pvntr\OneDrive\Desktop\Final 1\COMP-1682-001584160.docx'
d=Document(p)
terms=['stock transfer','shipping','React','Spring Boot','Java','month','Week','weeks','Implementation','Maintenance','inventory management']
for i,x in enumerate(d.paragraphs):
 t=' '.join(x.text.split())
 if any(term.lower() in t.lower() for term in terms):
  print(f'P{i}: {t}')
print('--- TABLES ---')
for ti,t in enumerate(d.tables,1):
 for ri,row in enumerate(t.rows):
  s=' | '.join(' '.join(c.text.split()) for c in row.cells)
  if any(term.lower() in s.lower() for term in terms): print(f'T{ti}R{ri}: {s}')
